"""
cerrar_inprg_antiguas.py
------------------------
Cierra OTs en estado INPRG con mas de N dias de antiguedad, sin importar
el tecnico asignado (WORKORDER.LEAD).

CONTEXTO:
    Limpieza de OTs antiguas en O_GESFO que han quedado abiertas en INPRG
    pero que en la practica deberian estar cerradas. Como el universo es
    grande y heterogeneo (multiples tecnicos asignados), las cerramos
    directamente via API en lugar de coordinar caso por caso.

FUENTE DE DATOS:
    Maximo MXAPIWODETAIL via integrations.maximo.rest_api.listar_ots().
    NO usa Postgres (la tabla onms.work_orders solo tiene OTs frescas,
    no el historico que necesitamos para este caso).

FLUJO DE ESTADOS EN MAXIMO:
    INPRG → COMP → CLOSE

    No se puede saltar directo de INPRG a CLOSE. Por eso el script aplica
    los DOS PATCH en secuencia (con pausa entre ellos para que Maximo
    refresque el statusdate).

UNIVERSO QUE PROCESA (default):
    - ownergroup = O_GESFO
    - classstructureid = 4213
    - status = INPRG
    - dias desde reportdate > DIAS_MINIMOS    ← default 90, configurable
    - sin filtro de tecnico                    ← TODAS las OTs candidatas

OPCIONALMENTE:
    --tecnicos T1 T2 ...    restringe el universo a esos tecnicos (campo lead)
    --excluir-tecnicos T1 T2 ...   excluye esos tecnicos del universo

MODOS DE EJECUCION:
    1. Por defecto: DRY-RUN. Lista las OTs candidatas y las exporta a Excel,
       NO toca Maximo. Sirve para revisar antes de ejecutar.
    2. Con --ejecutar: realmente hace los PATCH a Maximo (INPRG→COMP→CLOSE).
    3. Con --solo-comp: solo lleva a COMP y se detiene (mas conservador).

USO:
    # 1. Dry-run (SIEMPRE primero) - solo lista a Excel
    py -m etl.cerrar_inprg_antiguas

    # 2. Probar con un puñado primero antes del lote completo
    py -m etl.cerrar_inprg_antiguas --ejecutar --limit 3

    # 3. Ejecutar el cierre real (todas las candidatas)
    py -m etl.cerrar_inprg_antiguas --ejecutar

    # 4. Solo dejar en COMP (sin CLOSE) - mas reversible
    py -m etl.cerrar_inprg_antiguas --ejecutar --solo-comp

    # 5. Cambiar el umbral de dias
    py -m etl.cerrar_inprg_antiguas --dias 60

    # 6. Restringir a tecnicos especificos (campo lead)
    py -m etl.cerrar_inprg_antiguas --tecnicos LUIS.RODRIGUEZG NESTOR.RAMIREZ

    # 7. Excluir tecnicos (ej: dejar tranquilas las de LGMELENDEZHE)
    py -m etl.cerrar_inprg_antiguas --excluir-tecnicos LGMELENDEZHE

    # 8. Saltar el enriquecimiento (mas rapido, sin coordinador/lider/eecc)
    py -m etl.cerrar_inprg_antiguas --sin-enriquecer

SALIDA:
    output/Cierre_INPRG_Antiguas_{MODO}_{YYYYMMDD_HHMMSS}.xlsx
        Hoja 1 (Candidatas):  lista detallada de OTs a procesar
        Hoja 2 (Resultado):   resultado del cierre por OT (solo si --ejecutar)
"""

import argparse
import time
from collections import Counter
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from core.logging_setup import logger
from integrations.maximo.rest_api import (
    cambiar_estado,
    listar_ots,
    obtener_detalle_ot,
    obtener_ci_description,
)


# ════════════════════════════════════════════════════════════════════
# CONFIGURACION POR DEFECTO
# ════════════════════════════════════════════════════════════════════

DIAS_MINIMOS_DEFAULT = 90
OWNERGROUP           = "O_GESFO"
CLASSSTRUCTUREID     = "4213"

# Pausa entre PATCH INPRG→COMP y COMP→CLOSE. Maximo a veces necesita
# unos segundos para refrescar el statusdate antes del siguiente cambio.
PAUSA_ENTRE_PATCH_SEG = 2

# Pausa entre OTs (para no saturar Maximo)
PAUSA_ENTRE_OTS_SEG = 1

# Atributos del workorderspec que queremos extraer al enriquecer.
# El array workorderspec es una lista de objetos {assetattrid, alnvalue, ...}.
# Mapeamos assetattrid → clave de salida.
SPEC_ATTR_MAP = {
    "COORDINADOR_RED_FO":   "coordinador",
    "LIDER_DE_ZONA_FO":     "lider_zona",
    "EECC_CUADRILLA_FO":    "eecc",
    "NUMERO_CASO_FO":       "numero_caso",
    "TIPO_OPERACION_FO":    "tipo_operacion",
}


# ════════════════════════════════════════════════════════════════════
# COLUMNAS DEL EXCEL
# ════════════════════════════════════════════════════════════════════

COLUMNAS_CANDIDATAS = [
    ("wonum",            "WONUM",                 14),
    ("dias_abierta",     "Dias Abierta",          12),
    ("status",           "Status",                10),
    ("worktype",         "Worktype",              10),
    ("tipo_operacion",   "Tipo Operacion",        18),
    ("tecnico",          "Tecnico (lead)",        22),
    ("fecha_creacion",   "Fecha Creacion",        18),
    ("resumen",          "Resumen",               50),
    ("coordinador",      "Coordinador",           22),
    ("lider_zona",       "Lider Zona",            18),
    ("codigo_sitio",     "Cod. Sitio",            12),
    ("nombre_sitio",     "Nombre Sitio",          30),
    ("eecc",             "EECC",                  18),
    ("numero_caso",      "Numero Caso",           16),
    ("cinum",            "CI",                    30),
    ("ci_description",   "CI Descripcion",        40),
]

COLUMNAS_RESULTADO = [
    ("wonum",            "WONUM",                14),
    ("tecnico",          "Tecnico (lead)",       22),
    ("dias_abierta",     "Dias Abierta",         12),
    ("paso_comp_ok",     "INPRG→COMP OK",        14),
    ("paso_comp_msg",    "INPRG→COMP Mensaje",   40),
    ("paso_close_ok",    "COMP→CLOSE OK",        14),
    ("paso_close_msg",   "COMP→CLOSE Mensaje",   40),
    ("estado_final",     "Estado Final",         14),
    ("error",            "Error",                40),
]


# ════════════════════════════════════════════════════════════════════
# ESTILOS DE EXCEL
# ════════════════════════════════════════════════════════════════════

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
CELL_ALIGN = Alignment(horizontal="left", vertical="top", wrap_text=True)

THIN_BORDER = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)

FILL_OK      = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
FILL_ERROR   = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
FILL_PARCIAL = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid")


# ════════════════════════════════════════════════════════════════════
# UTILIDADES
# ════════════════════════════════════════════════════════════════════

def _parse_fecha_maximo(s):
    """
    Maximo devuelve fechas en ISO con offset, ej: '2025-08-12T14:23:00-05:00'.
    Devuelve SIEMPRE un datetime naive (sin tz) para poder restarlo con
    datetime.now(). Para nuestro caso (diff en dias) el offset no afecta
    porque tanto reportdate como now() estan en hora Colombia.
    """
    if not s:
        return None
    try:
        # Cortar a 19 chars (YYYY-MM-DDTHH:MM:SS) elimina cualquier offset/tz
        return datetime.strptime(s[:19], "%Y-%m-%dT%H:%M:%S")
    except Exception:
        return None


def _extraer_specs(detalle):
    """
    Recorre detalle['workorderspec'] y extrae los atributos definidos en
    SPEC_ATTR_MAP. Devuelve dict {clave_salida: valor}.
    Si una spec no aparece, su valor sera "".
    """
    out = {v: "" for v in SPEC_ATTR_MAP.values()}
    specs = detalle.get("workorderspec") or []
    for spec in specs:
        attrid = spec.get("assetattrid")
        if attrid in SPEC_ATTR_MAP:
            # Maximo guarda el valor en alnvalue, numvalue o tablevalue segun el tipo.
            # Para los campos que extraemos (todos texto) basta con alnvalue.
            valor = spec.get("alnvalue") or spec.get("numvalue") or spec.get("tablevalue") or ""
            out[SPEC_ATTR_MAP[attrid]] = valor
    return out


def _construir_candidata(member, dias_abierta, detalle=None, ci_cache=None):
    """
    Construye el dict de una candidata combinando los campos del listado
    (member) con el enriquecimiento del detalle (si se proveyo).
    """
    candidata = {
        "wonum":          member.get("wonum"),
        "status":         member.get("status"),
        "worktype":       member.get("worktype"),
        "resumen":        (member.get("description") or "").strip(),
        "codigo_sitio":   member.get("location"),
        "nombre_sitio":   member.get("nom_ubicacion"),
        "fecha_creacion": _parse_fecha_maximo(member.get("reportdate")),
        "cinum":          member.get("cinum"),
        "dias_abierta":   dias_abierta,
        # Campos que solo salen del detalle:
        "tecnico":        "",
        "coordinador":    "",
        "lider_zona":     "",
        "eecc":           "",
        "numero_caso":    "",
        "tipo_operacion": "",
        "ci_description": "",
    }

    if detalle is not None:
        # lead es un campo top-level del workorder
        candidata["tecnico"] = detalle.get("lead") or ""
        # Specs
        specs = _extraer_specs(detalle)
        candidata.update(specs)
        # CI description (top-level si viene; si no, consultamos)
        candidata["ci_description"] = (
            detalle.get("ci_description")
            or (obtener_ci_description(candidata["cinum"], cache=ci_cache) if candidata["cinum"] else "")
        )

    return candidata


# ════════════════════════════════════════════════════════════════════
# CIERRE DE UNA OT (INPRG → COMP → CLOSE)
# ════════════════════════════════════════════════════════════════════

def cerrar_ot_dos_pasos(wonum, solo_comp=False):
    """
    Lleva una OT de INPRG a CLOSE en dos PATCH:
        Paso 1: INPRG → COMP
        Paso 2: COMP → CLOSE  (omitido si solo_comp=True)
    """
    resultado = {
        "wonum":          wonum,
        "paso_comp_ok":   False,
        "paso_comp_msg":  "",
        "paso_close_ok":  False,
        "paso_close_msg": "",
        "estado_final":   "INPRG",
        "error":          "",
    }

    # PASO 1: INPRG → COMP
    try:
        r1 = cambiar_estado(wonum, "COMP")
        resultado["paso_comp_ok"]  = r1.get("success", False)
        resultado["paso_comp_msg"] = r1.get("message", "")
        if not resultado["paso_comp_ok"]:
            resultado["error"] = f"Fallo INPRG→COMP: {r1.get('message')}"
            return resultado
        resultado["estado_final"] = "COMP"
    except Exception as e:
        resultado["error"] = f"Excepcion en INPRG→COMP: {e}"
        return resultado

    if solo_comp:
        return resultado

    # PASO 2: COMP → CLOSE (con pausa)
    time.sleep(PAUSA_ENTRE_PATCH_SEG)

    try:
        r2 = cambiar_estado(wonum, "CLOSE")
        resultado["paso_close_ok"]  = r2.get("success", False)
        resultado["paso_close_msg"] = r2.get("message", "")
        if resultado["paso_close_ok"]:
            resultado["estado_final"] = "CLOSE"
        else:
            resultado["error"] = f"Fallo COMP→CLOSE: {r2.get('message')}"
    except Exception as e:
        resultado["error"] = f"Excepcion en COMP→CLOSE: {e}"

    return resultado


# ════════════════════════════════════════════════════════════════════
# CONSTRUCCION DE EXCEL
# ════════════════════════════════════════════════════════════════════

def aplicar_header(ws, columnas):
    for col_idx, (_, header, ancho) in enumerate(columnas, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = ancho
    ws.row_dimensions[1].height = 35
    ws.freeze_panes = "B2"


def construir_hoja_candidatas(ws, filas):
    aplicar_header(ws, COLUMNAS_CANDIDATAS)
    for row_idx, fila in enumerate(filas, start=2):
        for col_idx, (key, _, _) in enumerate(COLUMNAS_CANDIDATAS, start=1):
            valor = fila.get(key)
            if key == "dias_abierta" and valor is not None:
                valor = round(float(valor), 1)
            elif key == "fecha_creacion" and valor is not None:
                valor = valor.strftime("%Y-%m-%d %H:%M") if hasattr(valor, "strftime") else str(valor)
            cell = ws.cell(row=row_idx, column=col_idx, value=valor)
            cell.alignment = CELL_ALIGN
            cell.border = THIN_BORDER
    if filas:
        last_col = get_column_letter(len(COLUMNAS_CANDIDATAS))
        ws.auto_filter.ref = f"A1:{last_col}{len(filas) + 1}"


def construir_hoja_resultado(ws, resultados, solo_comp):
    aplicar_header(ws, COLUMNAS_RESULTADO)
    for row_idx, fila in enumerate(resultados, start=2):
        if solo_comp:
            ok = fila.get("paso_comp_ok")
        else:
            ok = fila.get("paso_comp_ok") and fila.get("paso_close_ok")
        parcial = fila.get("paso_comp_ok") and not fila.get("paso_close_ok") and not solo_comp

        if ok:
            fill = FILL_OK
        elif parcial:
            fill = FILL_PARCIAL
        else:
            fill = FILL_ERROR

        for col_idx, (key, _, _) in enumerate(COLUMNAS_RESULTADO, start=1):
            valor = fila.get(key)
            if isinstance(valor, bool):
                valor = "SI" if valor else "NO"
            elif key == "dias_abierta" and valor is not None:
                valor = round(float(valor), 1)
            cell = ws.cell(row=row_idx, column=col_idx, value=valor)
            cell.alignment = CELL_ALIGN
            cell.border = THIN_BORDER
            cell.fill = fill
    if resultados:
        last_col = get_column_letter(len(COLUMNAS_RESULTADO))
        ws.auto_filter.ref = f"A1:{last_col}{len(resultados) + 1}"


# ════════════════════════════════════════════════════════════════════
# MAIN
# ════════════════════════════════════════════════════════════════════

def main():
    parser = argparse.ArgumentParser(
        description="Cierra OTs INPRG antiguas en Maximo (INPRG → COMP → CLOSE)."
    )
    parser.add_argument(
        "--ejecutar",
        action="store_true",
        help="Realmente ejecuta el cierre en Maximo. Sin esta flag, solo dry-run.",
    )
    parser.add_argument(
        "--solo-comp",
        action="store_true",
        help="Solo lleva las OTs a COMP, no a CLOSE. Mas conservador.",
    )
    parser.add_argument(
        "--dias",
        type=int,
        default=DIAS_MINIMOS_DEFAULT,
        help=f"Dias minimos de antiguedad desde reportdate (default {DIAS_MINIMOS_DEFAULT}).",
    )
    parser.add_argument(
        "--tecnicos",
        nargs="+",
        default=None,
        help="Restringir el universo a estos tecnicos (campo lead). Si se omite, procesa TODOS.",
    )
    parser.add_argument(
        "--excluir-tecnicos",
        nargs="+",
        default=None,
        dest="excluir_tecnicos",
        help="Excluir estos tecnicos del universo. No combinable con --tecnicos.",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=None,
        help="Limite de OTs a procesar (util para pruebas controladas).",
    )
    parser.add_argument(
        "--sin-enriquecer",
        action="store_true",
        help="Salta la llamada a obtener_detalle_ot por OT. Mas rapido pero sin "
             "coordinador, lider_zona, eecc, numero_caso, tipo_operacion ni lead.",
    )
    args = parser.parse_args()

    if args.tecnicos and args.excluir_tecnicos:
        parser.error("--tecnicos y --excluir-tecnicos son mutuamente excluyentes.")

    # Si se pidio filtro de tecnicos hay que enriquecer si o si (lead solo
    # viene en el detalle, no en el listado paginado).
    if (args.tecnicos or args.excluir_tecnicos) and args.sin_enriquecer:
        parser.error("--tecnicos/--excluir-tecnicos requieren enriquecimiento; "
                     "no se puede combinar con --sin-enriquecer.")

    if args.tecnicos:
        filtro_tec = f"SOLO {args.tecnicos}"
    elif args.excluir_tecnicos:
        filtro_tec = f"TODOS excepto {args.excluir_tecnicos}"
    else:
        filtro_tec = "TODOS los tecnicos (sin filtro)"

    print("=" * 75)
    print("CIERRE DE OTs INPRG ANTIGUAS")
    print("=" * 75)
    print(f"  Modo:             {'EJECUCION REAL' if args.ejecutar else 'DRY-RUN (sin tocar Maximo)'}")
    print(f"  Estado meta:      {'COMP' if args.solo_comp else 'CLOSE'}")
    print(f"  Dias min:         {args.dias} (desde reportdate)")
    print(f"  Ownergroup:       {OWNERGROUP}")
    print(f"  Classstructureid: {CLASSSTRUCTUREID}")
    print(f"  Tecnicos:         {filtro_tec}")
    print(f"  Enriquecimiento:  {'NO (sin specs)' if args.sin_enriquecer else 'SI (detalle por OT)'}")
    print(f"  Limite:           {args.limit if args.limit else 'sin limite'}")
    print("=" * 75)

    # ════════════════════════════════════════════════════════
    # 1. LISTAR CANDIDATAS EN MAXIMO
    # ════════════════════════════════════════════════════════
    print("\n[1/4] Consultando OTs INPRG en Maximo (paginacion completa)...")

    # Filtro: reportdate < (hoy - dias). Asi Maximo devuelve solo OTs
    # con creacion mas antigua que el umbral.
    fecha_corte = datetime.now() - timedelta(days=args.dias)

    members = listar_ots(
        ownergroup       = OWNERGROUP,
        classstructureid = CLASSSTRUCTUREID,
        status_in        = ["INPRG"],
        fecha_hasta      = fecha_corte,   # reportdate < fecha_corte
    )

    if not members:
        print("\n[INFO] No hay OTs candidatas. Nada que hacer.")
        return

    print(f"      Total candidatas (pre-enriquecimiento): {len(members)}")

    # Cortar limite ANTES de enriquecer (para no llamar 26000 veces obtener_detalle_ot)
    if args.limit:
        members = members[: args.limit]
        print(f"      Recortado a --limit={args.limit}")

    # ════════════════════════════════════════════════════════
    # 2. ENRIQUECER CON DETALLE (workorderspec, lead, ci_desc)
    # ════════════════════════════════════════════════════════
    ahora = datetime.now()
    ci_cache = {}
    candidatas = []

    if args.sin_enriquecer:
        print("\n[2/4] Saltando enriquecimiento (--sin-enriquecer)...")
        for m in members:
            f_creacion = _parse_fecha_maximo(m.get("reportdate"))
            dias = (ahora - f_creacion).total_seconds() / 86400 if f_creacion else 0
            candidatas.append(_construir_candidata(m, dias))
    else:
        print(f"\n[2/4] Enriqueciendo {len(members)} OTs con detalle (workorderspec)...")
        for i, m in enumerate(members, 1):
            wonum = m.get("wonum")
            href  = m.get("href")

            f_creacion = _parse_fecha_maximo(m.get("reportdate"))
            dias = (ahora - f_creacion).total_seconds() / 86400 if f_creacion else 0

            detalle = obtener_detalle_ot(href) if href else None
            cand = _construir_candidata(m, dias, detalle=detalle, ci_cache=ci_cache)
            candidatas.append(cand)

            if i % 25 == 0 or i == len(members):
                print(f"      Enriquecidas: {i}/{len(members)}")

    # ════════════════════════════════════════════════════════
    # 2.5 FILTRO POR TECNICO (post-enriquecimiento)
    # ════════════════════════════════════════════════════════
    if args.tecnicos:
        tecnicos_set = set(args.tecnicos)
        antes = len(candidatas)
        candidatas = [c for c in candidatas if c.get("tecnico") in tecnicos_set]
        print(f"      Filtro --tecnicos: {antes} → {len(candidatas)}")
    elif args.excluir_tecnicos:
        excluir_set = set(args.excluir_tecnicos)
        antes = len(candidatas)
        candidatas = [c for c in candidatas if c.get("tecnico") not in excluir_set]
        print(f"      Filtro --excluir-tecnicos: {antes} → {len(candidatas)}")

    if not candidatas:
        print("\n[INFO] No hay OTs tras aplicar filtros. Nada que hacer.")
        return

    # Resumen por tecnico
    por_tecnico = Counter((c.get("tecnico") or "(sin tecnico)") for c in candidatas)
    print(f"      Distribucion por tecnico (top 15):")
    for tec, cant in por_tecnico.most_common(15):
        print(f"        {tec:<25} {cant:>5}")
    if len(por_tecnico) > 15:
        print(f"        ... y {len(por_tecnico) - 15} tecnicos mas")

    # ════════════════════════════════════════════════════════
    # 3. PROCESAR CIERRE (si --ejecutar)
    # ════════════════════════════════════════════════════════
    resultados = []
    if args.ejecutar:
        print(f"\n[3/4] Procesando {len(candidatas)} OTs en Maximo...")
        print(f"      (pausa de {PAUSA_ENTRE_PATCH_SEG}s entre PATCH y {PAUSA_ENTRE_OTS_SEG}s entre OTs)")

        for i, c in enumerate(candidatas, 1):
            wonum   = c["wonum"]
            tecnico = c.get("tecnico") or "(sin lead)"
            dias    = c["dias_abierta"]

            print(f"  [{i}/{len(candidatas)}] OT {wonum} (lead={tecnico}, dias={float(dias):.1f}) ... ",
                  end="", flush=True)

            res = cerrar_ot_dos_pasos(wonum, solo_comp=args.solo_comp)
            res["tecnico"]      = tecnico
            res["dias_abierta"] = dias
            resultados.append(res)

            if args.solo_comp:
                if res["paso_comp_ok"]:
                    print(f"OK → {res['estado_final']}")
                else:
                    print(f"FAIL ({res['paso_comp_msg']})")
            else:
                if res["paso_comp_ok"] and res["paso_close_ok"]:
                    print(f"OK → CLOSE")
                elif res["paso_comp_ok"]:
                    print(f"PARCIAL → COMP (CLOSE fallo: {res['paso_close_msg']})")
                else:
                    print(f"FAIL ({res['paso_comp_msg']})")

            if i < len(candidatas):
                time.sleep(PAUSA_ENTRE_OTS_SEG)
    else:
        print("\n[3/4] DRY-RUN: omitiendo ejecucion en Maximo.")

    # ════════════════════════════════════════════════════════
    # 4. GENERAR EXCEL
    # ════════════════════════════════════════════════════════
    print(f"\n[4/4] Generando Excel...")
    wb = Workbook()

    ws_cand = wb.active
    ws_cand.title = "Candidatas"
    construir_hoja_candidatas(ws_cand, candidatas)

    if resultados:
        ws_res = wb.create_sheet(title="Resultado")
        construir_hoja_resultado(ws_res, resultados, args.solo_comp)

    output_dir = Path("output")
    output_dir.mkdir(exist_ok=True)

    sufijo = "EJECUCION" if args.ejecutar else "DRYRUN"
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = output_dir / f"Cierre_INPRG_Antiguas_{sufijo}_{timestamp}.xlsx"
    wb.save(filename)

    # Resumen final
    print(f"\n{'=' * 75}")
    print(f"RESUMEN")
    print(f"{'=' * 75}")
    print(f"  Candidatas:             {len(candidatas)}")
    if resultados:
        ok_total = sum(
            1 for r in resultados
            if (r["paso_comp_ok"] and r["paso_close_ok"]) or
               (args.solo_comp and r["paso_comp_ok"])
        )
        parciales = sum(
            1 for r in resultados
            if r["paso_comp_ok"] and not r["paso_close_ok"] and not args.solo_comp
        )
        fallos = len(resultados) - ok_total - parciales
        print(f"  OK total:               {ok_total}")
        if not args.solo_comp:
            print(f"  Parciales (COMP solo):  {parciales}")
        print(f"  Fallos:                 {fallos}")
    print(f"  Excel:                  {filename}")
    print(f"{'=' * 75}")


if __name__ == "__main__":
    main()